#!/usr/bin/env python3
"""Build DESIGN.pptx from DESIGN.md — one slide per numbered section, with the
speaker notes of each slide carrying a verbatim copy of that section's markdown.

    pip install python-pptx
    python3 tools/build-design-deck.py

Re-run after editing DESIGN.md: the notes are read from the file, so they cannot
drift from the doc. The slide bodies are hand-authored here — when a section's
substance changes, update the corresponding block below too.

Colours are the validated categorical palette (blue / orange / aqua) plus the
status set; the one quantitative mark (the bundle-composition bar on the §5
slide) carries direct labels on every segment.

Note: Keynote refuses to import the generated file ("format is invalid") even
though the package is valid OOXML — open it with PowerPoint.
"""
import os
import re
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # the cp4d/ dir
SRC = os.path.join(HERE, "DESIGN.md")
OUT = os.path.join(HERE, "DESIGN.pptx")

# --- palette (validated: slots 1-3 pass all-pairs light on #ffffff) ----------
BLUE   = C(0x2A, 0x78, 0xD6)   # categorical slot 1
ORANGE = C(0xEB, 0x68, 0x34)   # slot 2
AQUA   = C(0x1B, 0xAF, 0x7A)   # slot 3
CRIT   = C(0xD0, 0x3B, 0x3B)   # status critical
GOOD   = C(0x0C, 0xA3, 0x0C)   # status good
WARN   = C(0xFA, 0xB2, 0x19)   # status warning
INK    = C(0x0B, 0x0B, 0x0B)   # primary ink
INK2   = C(0x52, 0x51, 0x4E)   # secondary ink
MUTED  = C(0x89, 0x87, 0x81)   # muted
HAIR   = C(0xE1, 0xE0, 0xD9)   # gridline hairline
RULE   = C(0xC3, 0xC2, 0xBB)   # baseline
SURF   = C(0xFC, 0xFC, 0xFB)   # chart surface
PLANE  = C(0xF4, 0xF4, 0xF1)   # recessive panel
WHITE  = C(0xFF, 0xFF, 0xFF)
FONT   = "Calibri"

W, H = In(13.333), In(7.5)
M = In(0.62)                    # page margin
CW = W - 2 * M                  # content width


# --- primitives --------------------------------------------------------------
def txbox(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False, space_after=4, line=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, raw in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line:
            p.line_spacing = line
        # inline **bold** runs
        for j, chunk in enumerate(re.split(r"(\*\*[^*]+\*\*)", raw)):
            if not chunk:
                continue
            r = p.add_run()
            strong = chunk.startswith("**") and chunk.endswith("**")
            r.text = chunk[2:-2] if strong else chunk
            f = r.font
            f.name, f.size, f.color.rgb = FONT, Pt(size), color
            f.bold = bold or strong
            f.italic = italic
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0, dash=None):
    s = slide.shapes.add_shape(kind, x, y, w, h)
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = dash
    s.text_frame.word_wrap = True
    return s


def box(slide, x, y, w, h, title=None, body=None, fill=SURF, line=HAIR, accent=None,
        title_size=15, body_size=12, title_color=INK, body_color=INK2, pad=In(0.16),
        title_h=In(0.34)):
    """Rounded panel, optional 4px accent stripe down the left edge."""
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line)
    s.adjustments[0] = 0.06
    if accent:
        shape(slide, MSO_SHAPE.RECTANGLE, x, y + In(0.1), Pt(4), h - In(0.2), fill=accent)
    ty = y + pad
    if title:
        txbox(slide, x + pad + (In(0.1) if accent else 0), ty, w - 2 * pad, title_h,
              title, size=title_size, bold=True, color=title_color)
        ty += title_h
    if body:
        txbox(slide, x + pad + (In(0.1) if accent else 0), ty, w - 2 * pad, h - (ty - y) - pad,
              body, size=body_size, color=body_color, line=1.15, space_after=5)
    return s


def chip(slide, x, y, w, text, fill=None, line=RULE, color=INK2, size=11, h=In(0.34), bold=False):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line)
    s.adjustments[0] = 0.5
    tf = s.text_frame
    tf.margin_left = tf.margin_right = In(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.color.rgb, r.font.bold = FONT, Pt(size), color, bold
    return s


def arrow(slide, x, y, w, h, fill=RULE, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, x, y, w, h, fill=fill)


def hrule(slide, x, y, w, color=HAIR, pt=1.0):
    ln = slide.shapes.add_connector(1, x, y, x + w, y)
    ln.line.color.rgb = color
    ln.line.width = Pt(pt)
    return ln


# --- slide frame -------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
SECTIONS = {}      # number -> markdown source, filled below
DECK = []          # (slide, notes_key)


def new_slide(kicker, title, notes):
    s = prs.slides.add_slide(BLANK)
    txbox(s, M, In(0.42), CW, In(0.2), kicker.upper(), size=11, bold=True, color=BLUE)
    txbox(s, M, In(0.66), CW, In(0.42), title, size=27, bold=True, color=INK)
    hrule(s, M, In(1.24), CW, RULE)
    txbox(s, M, H - In(0.44), CW, In(0.2),
          "CP4D artifact-level backup with Kasten · DESIGN.md", size=9, color=MUTED)
    s.notes_slide.notes_text_frame.text = notes
    return s


# --- parse DESIGN.md into sections -------------------------------------------
md = open(SRC).read()
intro, *rest = re.split(r"\n(?=## )", md)
for chunk in rest:
    n = re.match(r"## (\d+)\.", chunk)
    if n:
        SECTIONS[int(n.group(1))] = chunk.strip()
INTRO = intro.strip()
assert set(SECTIONS) == set(range(1, 8)), SECTIONS.keys()

Y0 = In(1.52)      # top of content band


# ============================ 0 · title ======================================
s = prs.slides.add_slide(BLANK)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, Pt(6), H, fill=BLUE)
txbox(s, In(1.1), In(2.1), In(11), In(0.24), "CLOUD PAK FOR DATA · KASTEN BLUEPRINT",
      size=12, bold=True, color=BLUE)
txbox(s, In(1.1), In(2.5), In(11), In(1.0), "Artifact-level backup — the design",
      size=40, bold=True, color=INK)
txbox(s, In(1.1), In(3.6), In(9.6), In(1.3),
      "Why the API is the only path to a notebook, how it maps onto the Kasten blueprint "
      "patterns, and what the export bundle actually contains.",
      size=16, color=INK2, line=1.25)
hrule(s, In(1.1), In(5.05), In(6.2), RULE)
txbox(s, In(1.1), In(5.25), In(11), In(0.9),
      "Validated end-to-end · Kasten 8.5.12 · OpenShift 4.18.6 · CP4D 5.2.x · cpdctl 1.8.244\n"
      "Speaker notes on every slide carry the full text of the matching DESIGN.md section.",
      size=12, color=MUTED, line=1.3)
s.notes_slide.notes_text_frame.text = INTRO

# ============================ 1 · where this fits ============================
s = new_slide("Section 1", "Where this fits", SECTIONS[1])
half = (CW - In(0.34)) / 2
box(s, M, Y0, half, In(2.96), "cpdbr / cpd-cli oadp",
    "Platform-level backup that ships with CP4D.\n\n"
    "•  Scope: the whole tenant — all of CP4D at once\n"
    "•  Coupled to a resource-backup engine plus a separate object store\n"
    "•  Built for infrastructure DR of the entire installation\n"
    "•  Protects the control plane",
    fill=PLANE, accent=MUTED, body_size=13)
box(s, M + half + In(0.34), Y0, half, In(2.96), "This blueprint",
    "Artifact-level backup of analytics projects.\n\n"
    "•  Scope: one project at a time\n"
    "•  Granular, self-service restore\n"
    "•  Portability — migrate a project to another cluster\n"
    "•  Needs the platform up; does not protect the control plane",
    fill=SURF, accent=BLUE, body_size=13)
box(s, M, Y0 + In(3.24), CW, In(1.25), "A complement, not a replacement",
    "Two different jobs, not two options for the same job. Platform DR restores the installation; "
    "this restores a project. Positioning it as a DR replacement is the one mistake to avoid — "
    "it needs the CP4D platform to be up, and it does not protect the control plane.",
    fill=SURF, accent=ORANGE, body_size=13)

# ============================ 2 · why the API ================================
s = new_slide("Section 2", "What a CP4D artifact is — and why the API is the only path",
              SECTIONS[2])
lw = CW - In(1.9)                      # layer stack width; door box to its right
cw5 = (lw - In(0.56) - In(4 * 0.14)) / 5   # chip width inside a layer
# application layer
box(s, M, Y0, lw, In(1.42), "CP4D application layer",
    "Entities in CP4D's metadata database + object storage.", fill=SURF, accent=BLUE,
    body_size=12)
for i, t in enumerate(["Project", "Notebook", "Job", "Connection", "Data asset"]):
    chip(s, M + In(0.28) + i * (cw5 + In(0.14)), Y0 + In(0.92), cw5, t, fill=WHITE, line=BLUE,
         color=BLUE, size=11)
# the wall
wall_y = Y0 + In(1.62)
ln = s.shapes.add_connector(1, M, wall_y, M + lw, wall_y)
ln.line.color.rgb = CRIT
ln.line.width = Pt(2.5)
ln.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
txbox(s, M, wall_y + In(0.07), lw, In(0.24),
      "`oc get` will never show a notebook — no CRD, no PVC, no ConfigMap",
      size=11, bold=True, color=CRIT, align=PP_ALIGN.CENTER)
# kubernetes layer
ky = wall_y + In(0.4)
box(s, M, ky, lw, In(1.42), "Kubernetes layer",
    "Everything a resource backup can see.", fill=PLANE, accent=MUTED, body_size=12)
for i, t in enumerate(["Namespaces", "Pods", "PVCs", "CRs", "Secrets"]):
    chip(s, M + In(0.28) + i * (cw5 + In(0.14)), ky + In(0.92), cw5, t, fill=WHITE,
         line=RULE, color=INK2, size=11)
# the only door
door_x = M + lw + In(0.34)
door_w = CW - lw - In(0.34)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, door_x, Y0, door_w, In(1.42), fill=BLUE)
txbox(s, door_x, Y0 + In(0.44), door_w, In(0.6), "cpdctl\nREST API", size=15, bold=True,
      color=WHITE, align=PP_ALIGN.CENTER)
txbox(s, door_x, Y0 + In(1.5), door_w, In(0.24), "the only door in", size=10.5, bold=True,
      color=BLUE, align=PP_ALIGN.CENTER)
arrow(s, door_x - In(0.34), Y0 + In(0.6), In(0.34), In(0.2), fill=BLUE,
      kind=MSO_SHAPE.LEFT_ARROW)
# three facts, full width along the bottom
fy = ky + In(1.66)
facts = [("A notebook is not a file, and not a CRD",
          "A Kubernetes resource backup cannot see it at all — so the API path is not the "
          "cleaner option here, it is the only one."),
         ("The unit is the project, not the notebook",
          "A notebook without its runtime binding, data assets and relationship graph is not "
          "usefully restorable."),
         ("The tool is cpdctl, not cpd-cli",
          "cpd-cli is the platform/install CLI — the cpdbr world. cpdctl is the runtime CLI for "
          "projects and assets.")]
fwd = (CW - In(2 * 0.26)) / 3
for i, (t, sub) in enumerate(facts):
    box(s, M + i * (fwd + In(0.26)), fy, fwd, In(1.28), None,
        "**" + str(i + 1) + ".  " + t + "**\n" + sub, fill=SURF, accent=AQUA, body_size=11.5,
        pad=In(0.13))
txbox(s, M, fy + In(1.46), CW, In(0.34),
      "Consequence: **PVC = unit of granularity = unit of multi-tenant restore.** "
      "One project, one PVC.", size=13, color=INK)

# ============================ 3 · pattern mapping ============================
s = new_slide("Section 3", "How it maps onto the blueprint patterns", SECTIONS[3])
txbox(s, M, Y0, CW, In(0.3),
      "**Pattern 4 (dump to a permanent keeper PVC) + the action-hook mechanism.** "
      "Kasten is the data mover; the blueprint never transfers data.",
      size=13, color=INK2)
# ordering timeline
ty = Y0 + In(0.5)
txbox(s, M, ty, CW, In(0.22), "WHY AN ACTION HOOK — DISCOVERY ORDERING", size=10, bold=True,
      color=MUTED)
tl = ty + In(0.34)
seg = In(2.62)
labels = [("BackupAction\npreHook", BLUE), ("PVC\ndiscovery", MUTED),
          ("CSI\nsnapshots", MUTED), ("restore point", MUTED)]
for i, (t, col) in enumerate(labels):
    chip(s, M + i * (seg + In(0.28)), tl, seg, t, fill=WHITE,
         line=col, color=col if col is BLUE else INK2, size=11, h=In(0.56),
         bold=col is BLUE)
    if i < 3:
        arrow(s, M + i * (seg + In(0.28)) + seg + In(0.03), tl + In(0.19), In(0.22), In(0.18))
txbox(s, M, tl + In(0.66), CW, In(0.5),
      "The set of projects — and therefore the set of PVCs — is only known at run time. "
      "A per-project PVC created in the preHook lands in the restore point; the same PVC created "
      "in a resource-bound `backupPrehook` would be created **after** discovery and silently "
      "left out. Trade-off: namespace-only context, one blueprint per namespace.",
      size=12, color=INK2, line=1.15)
# flow
fy = tl + In(1.36)
txbox(s, M, fy, CW, In(0.22), "WHAT THE HOOK DOES", size=10, bold=True, color=MUTED)
fy += In(0.32)
steps = ["enumerate\nprojects", "ensure one PVC\nper project", "launch one\nexport pod each",
         "prune deleted\nprojects", "delete export\npods"]
sw = (CW - In(4 * 0.26)) / 5
for i, t in enumerate(steps):
    chip(s, M + i * (sw + In(0.26)), fy, sw, t, fill=PLANE, line=HAIR, color=INK2,
         size=11, h=In(0.56))
    if i < 4:
        arrow(s, M + i * (sw + In(0.26)) + sw + In(0.02), fy + In(0.19), In(0.22), In(0.18))
txbox(s, M, fy + In(0.66), CW, In(0.3),
      "→ then **Kasten snapshots those PVCs**. No KanisterBackupData, no KanisterRestoreData.",
      size=12, color=INK)
# decisions
dy = fy + In(1.06)
txbox(s, M, dy, CW, In(0.22), "DECISIONS THAT FOLLOW", size=10, bold=True, color=MUTED)
dy += In(0.3)
dec = [("Permanent GUID-keyed PVCs", "so CSI block-level dedup works across runs"),
       ("Store the export UNZIPPED", "a zip is opaque to dedup; one cell edit = one file"),
       ("Restore decoupled from Kasten hooks", "a standalone script fits multi-tenant restore"),
       ("Credentials outside the backed-up ns", "never captured in any restore point")]
dwe = (CW - In(3 * 0.2)) / 4
for i, (t, sub) in enumerate(dec):
    x = M + i * (dwe + In(0.2))
    box(s, x, dy, dwe, In(0.86), None, "**" + t + "**\n" + sub, fill=SURF, accent=AQUA,
        body_size=10.5, pad=In(0.1))

# ============================ 4 · scope ======================================
s = new_slide("Section 4", "V1 scope, and the boundaries drawn on purpose", SECTIONS[4])
lwid = In(4.3)
box(s, M, Y0, lwid, In(2.5), "In scope",
    "•  Notebooks\n•  Jobs\n•  File data assets\n•  Connection definitions\n"
    "•  The dependency graph\n•  Environment / runtime definitions",
    fill=SURF, accent=GOOD, body_size=13)
rx = M + lwid + In(0.3)
rwid = CW - lwid - In(0.3)
box(s, rx, Y0, rwid, In(2.5), "Deliberately deferred",
    "**Custom runtime images** — export captures the environment definition, not the image in "
    "the internal registry. Untested.\n"
    "**External data behind connections** — reference only (§6); separately protected.\n"
    "**Version skew** — logical export/import is version-sensitive; V1 targets 5.2 → 5.2 (§7).\n"
    "**Git-based projects** — validated only against COS-backed (`assetfiles`).",
    fill=PLANE, accent=WARN, body_size=12)
box(s, M, Y0 + In(2.72), CW, In(0.72), None,
    "**Not a deferral — a requirement.** Connection credentials were assumed to be redacted on "
    "export. They are not (§6). That finding changed the design.",
    fill=SURF, accent=CRIT, body_size=12.5, pad=In(0.13))
box(s, M, Y0 + In(3.58), CW, In(1.76), "Git-based does not remove the need to back up the metadata",
    "**Git is tamper-evident, not tamper-proof.** Hashes and signatures let you detect a "
    "force-push, a rewritten branch or a deleted repo — they do not give you the previous state "
    "back. Detection is not recovery.\n"
    "**Git is usually inside the blast radius.** A disaster or ransomware event takes the git "
    "server, its runners and the credentials that reach them.\n"
    "→ **Two things to protect, not one:** this blueprint for the CP4D metadata, an independent "
    "process for the git remote. Git-based projects are not simpler.",
    fill=SURF, accent=CRIT, body_size=12)

# ============================ 5 · what export produces ======================
s = new_slide("Section 5", "What export/import actually produces", SECTIONS[5])
txbox(s, M, Y0, CW, In(0.3),
      "Fixture: one project, a 2-cell notebook + the job that runs it. Bundle: **147 KB, 76 files** "
      "— for a **995-byte** notebook.", size=13, color=INK2)
# 100% stacked bar (2 segments, direct-labelled — relief for the WARN'd hues)
by = Y0 + In(0.62)
bh = In(0.62)
bw = CW
GAP = Pt(2)
seg1 = Emu(int(bw * 0.70)) - GAP        # assettypes schemas
seg2 = Emu(int(bw * 0.30))
s1 = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M, by, seg1, bh, fill=MUTED)
s1.adjustments[0] = 0.10
s2 = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, M + seg1 + GAP, by, seg2, bh, fill=BLUE)
s2.adjustments[0] = 0.22
txbox(s, M + In(0.14), by + In(0.16), seg1 - In(0.28), In(0.3),
      "assettypes/ — 70 asset-type schemas (the metamodel)   ~70%", size=12.5, bold=True,
      color=WHITE)
txbox(s, M + seg1 + In(0.2), by + In(0.16), seg2 - In(0.3), In(0.3),
      "the payload   ~30%", size=12.5, bold=True, color=WHITE)
txbox(s, M, by + bh + In(0.12), CW, In(0.6),
      "Schema overhead is roughly **constant regardless of content** — which is exactly why the "
      "bundle is stored **unzipped**: a one-cell edit rewrites one file, and the constant schema "
      "mass dedups across snapshots instead of being re-shipped inside a fresh archive.",
      size=12, color=INK2, line=1.15)
# the three layers
ly = by + bh + In(0.8)
lay = [("Artifacts — the payload, and it is tiny", BLUE,
        "`assets/notebook/*.ipynb` is the real notebook, cells and outputs included. The job "
        "asset was captured too — `--assets-all-assets` really means all project assets."),
       ("Graph + project definition — small", AQUA,
        "`assetrelationships.json` records the job→notebook \"uses\" relationship. `project.json` "
        "shows `storage.type = assetfiles` — a COS-backed, not git-based, project."),
       ("assettypes/ — 70 files, ~70%", MUTED,
        "Asset-type schemas — the metamodel, not instances. Why a 995-byte notebook yields a "
        "147 KB bundle.")]
lwd = (CW - In(2 * 0.24)) / 3
for i, (t, col, sub) in enumerate(lay):
    box(s, M + i * (lwd + In(0.24)), ly, lwd, In(1.32), None, "**" + t + "**\n" + sub,
        fill=SURF, accent=col, body_size=11, pad=In(0.12))
# round trip
ry = ly + In(1.5)
txbox(s, M, ry, CW, In(0.24), "ROUND-TRIP FIDELITY — VERIFIED", size=10, bold=True, color=MUTED)
rows = [("Notebook content (cells + outputs)", "byte-for-byte identical"),
        ("Assets in the target project", "notebook + job, both available, fresh IDs"),
        ("Relationship graph", "job asset_ref → the new notebook ID"),
        ("Stock runtime", "job env_id → rt241py-<new-project-id>")]
cwid = (CW - In(3 * 0.2)) / 4
for i, (k, v) in enumerate(rows):
    x = M + i * (cwid + In(0.2))
    box(s, x, ry + In(0.3), cwid, In(0.66), None, "✓  **" + k + "**\n" + v, fill=PLANE,
        line=HAIR, body_size=10.5, pad=In(0.1))
txbox(s, M, ry + In(1.06), CW, In(0.4),
      "**Nuance:** on import the notebook loses its explicit `runtime.environment`, but the job "
      "keeps its runtime binding, remapped to the target project's stock runtime. Reproducibility "
      "rides on the **job**. Custom runtimes: still untested.",
      size=11.5, color=INK2, line=1.15)

# ============================ 6 · the boundary ===============================
s = new_slide("Section 6", "The data and credential boundary", SECTIONS[6])
bwid = In(6.5)
box(s, M, Y0, bwid, In(2.72), "What the export bundle contains", None, fill=SURF, accent=BLUE)
items = [("✓", GOOD, "Notebook `.ipynb`", "cells + outputs"),
         ("✓", GOOD, "File data asset", "**bytes included**"),
         ("✓", GOOD, "Connection definition", "`assets/.METADATA/connection.*.json`"),
         ("⚠", WARN, "Connected data asset", "**reference only** — `is_remote:true`, `size:0`")]
for i, (mark, col, k, v) in enumerate(items):
    iy = Y0 + In(0.62) + i * In(0.5)
    txbox(s, M + In(0.2), iy, In(0.3), In(0.3), mark, size=14, bold=True, color=col)
    txbox(s, M + In(0.52), iy + In(0.02), bwid - In(0.72), In(0.3),
          "**" + k + "** — " + v, size=12, color=INK2)
# external datasource
ex = M + bwid + In(0.42)
box(s, ex, Y0, CW - bwid - In(0.42), In(2.72), "Outside the bundle",
    "The S3 objects.\nThe rows in a Db2.\nThe files on an on-prem share.\n\n"
    "A connected data asset travels as a pointer — the bytes never leave the datasource.",
    fill=PLANE, accent=MUTED, body_size=13)
arrow(s, M + bwid + In(0.06), Y0 + In(1.9), In(0.32), In(0.2), fill=WARN)
# cleartext callout
cy = Y0 + In(2.94)
box(s, M, cy, CW, In(1.12), "Connection credentials are exported in CLEARTEXT",
    "Not redacted by default. `--encryption-key` on export encrypts the masked properties; "
    "**import needs the same key**, and losing the key means losing the credentials in the "
    "restored connections. The bundle is sensitive either way — lock down the PVCs and the "
    "namespace regardless.",
    fill=SURF, accent=CRIT, title_color=CRIT, body_size=12.5)
box(s, M, cy + In(1.26), CW, In(1.06), "What cpdctl does not protect",
    "Export captures the project, not the external data behind a connection — a deliberate "
    "boundary, and the same one cpdbr has. Per connection: no action if the datasource is "
    "versioned or a system-of-record with its own backups; otherwise a separate, independent "
    "protection process.",
    fill=SURF, accent=AQUA, body_size=12)

# ============================ 7 · version policy =============================
s = new_slide("Section 7", "Why “latest cpdctl” is the right version policy", SECTIONS[7])
txbox(s, M, Y0, CW, In(0.3),
      "Install the **newest stable cpdctl**, not one matched to your CP4D release — `cpdctl` "
      "versioning (1.8.x) is decoupled from CP4D versioning (4.x / 5.x).", size=13, color=INK2)
# client -> backends
gy = Y0 + In(0.56)
box(s, M, gy, In(2.5), In(1.7), None, "latest stable\n**cpdctl**\n1.8.244", fill=BLUE,
    line=None, body_color=WHITE, body_size=15)
for i, (ver, ok) in enumerate([("CP4D 5.2.x", True), ("CP4D 5.3.x", True),
                               ("any supported CP4D", True), ("long-EOL CP4D", False)]):
    yy = gy + i * In(0.44)
    arrow(s, M + In(2.6), yy + In(0.1), In(0.5), In(0.16),
          fill=GOOD if ok else RULE)
    chip(s, M + In(3.2), yy, In(2.5), ver, fill=WHITE, line=GOOD if ok else RULE,
         color=INK2, size=11, h=In(0.36))
    txbox(s, M + In(5.82), yy + In(0.06), In(1.6), In(0.24),
          "✓ works" if ok else "⚠ may fall off", size=11, bold=True,
          color=GOOD if ok else MUTED)
box(s, M + In(7.5), gy, CW - In(7.5), In(1.7), "Caveat 1 — “supported” is load-bearing",
    "“Backward compatible with all supported Cloud Pak for Data releases” is scoped to "
    "CP4D versions still in IBM support. A long-EOL backend can fall outside it.",
    fill=SURF, accent=WARN, body_size=12)
# bundle portability
py = gy + In(2.0)
txbox(s, M, py, CW, In(0.24),
      "CAVEAT 2 — CLIENT↔BACKEND COMPATIBILITY IS NOT BUNDLE↔BUNDLE PORTABILITY",
      size=10, bold=True, color=MUTED)
py += In(0.34)
chip(s, M, py, In(2.9), "bundle exported\nagainst CP4D 5.2.x", fill=PLANE, line=HAIR,
     color=INK2, size=11, h=In(0.62))
arrow(s, M + In(3.0), py + In(0.22), In(0.6), In(0.18), fill=WARN)
chip(s, M + In(3.7), py, In(2.9), "imported into\nCP4D 5.3.x", fill=PLANE, line=HAIR,
     color=INK2, size=11, h=In(0.62))
txbox(s, M + In(6.75), py + In(0.14), In(1.2), In(0.34), "not promised", size=13, bold=True,
      color=CRIT)
box(s, M + In(8.2), py - In(0.06), CW - In(8.2), In(0.76), None,
    "The bundle format is a property of the **CP4D backend that produced it**, not of the CLI.",
    fill=SURF, accent=CRIT, body_size=12, pad=In(0.12))
box(s, M, py + In(0.86), CW, In(0.78), None,
    "**So:** “latest cpdctl handles all CP4D versions” means the CLI can *talk to* them. "
    "V1 targets same-version DR (5.2 → 5.2). Cross-version migration is a separate validation — "
    "it is the version-skew risk listed in §4.",
    fill=PLANE, line=HAIR, body_size=12.5, pad=In(0.13))

prs.save(OUT)
print("wrote", OUT, "·", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
